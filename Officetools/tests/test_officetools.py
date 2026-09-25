"""核心层单测：bootstrap 自举 / parsers 五格式 / worker 入站流程 / S1 安全校验。

依赖：CABCLAW_HOST 宿主注入 + pylibs（conftest 自举）；样例由 fixtures 现场生成。
"""
from __future__ import annotations

import json
import os
import subprocess
import sys
import time
from pathlib import Path

import pytest

MODULE_DIR = Path(__file__).resolve().parent.parent
import bootstrap  # noqa: E402
import parsers  # noqa: E402
import officetools_worker as worker  # noqa: E402


# ---------- fixtures：现场生成样例 ----------

@pytest.fixture(scope="session")
def samples(tmp_path_factory) -> dict[str, Path]:
    out = tmp_path_factory.mktemp("samples")
    made: dict[str, Path] = {}

    import docx
    d = docx.Document()
    d.add_heading("项目计划", level=1)
    d.add_paragraph("这是第一段介绍文字。")
    d.add_heading("里程碑", level=2)
    d.add_paragraph("三月完成开发。")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "阶段"
    t.cell(0, 1).text = "时间"
    t.cell(1, 0).text = "开发"
    t.cell(1, 1).text = "3 月"
    p = out / "sample.docx"
    d.save(str(p))
    made["docx"] = p

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "预算"
    ws.append(["项目", "金额"])
    ws.append(["服务器", 5000])
    wb.save(str(out / "sample.xlsx"))
    made["xlsx"] = out / "sample.xlsx"

    from pptx import Presentation
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "季度汇报"
    s1.placeholders[1].text = "2026 Q3"
    prs.save(str(out / "sample.pptx"))
    made["pptx"] = out / "sample.pptx"

    pdf = b"""%PDF-1.4
1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj
2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj
3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>endobj
4 0 obj<</Length 90>>stream
BT /F1 12 Tf 72 720 Td (Hello Officetools PDF test document.) Tj ET
BT /F1 10 Tf 72 700 Td (Second line of body text here.) Tj ET
endstream
endobj
5 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj
trailer<</Root 1 0 R>>
"""
    p = out / "sample.pdf"
    p.write_bytes(pdf)
    made["pdf"] = p
    return made


@pytest.fixture()
def env(tmp_path, monkeypatch):
    """隔离环境：独立 inbox / 数据区 / settings。"""
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    data = tmp_path / "data"
    monkeypatch.setattr(worker, "inbox_dir", lambda: inbox)
    monkeypatch.setattr(worker, "DATA_DIR", data)
    monkeypatch.setattr(worker, "STATE_FILE", data / "state.json")
    monkeypatch.setattr(worker, "OUTPUTS_DIR", data / "outputs")
    monkeypatch.setattr(worker, "LOCK_FILE", data / ".parse.lock")
    monkeypatch.setattr(worker, "_settings", lambda: dict(worker.DEFAULT_SETTINGS))
    return {"inbox": inbox, "data": data}


# ---------- bootstrap ----------

class TestBootstrap:
    def test_ensure_core_ready(self):
        ok, detail = bootstrap.ensure_core()
        assert ok, detail

    def test_pylibs_env_override(self, tmp_path, monkeypatch):
        monkeypatch.setenv("OFFICETOOLS_PYLIBS", str(tmp_path / "pl"))
        assert bootstrap.pylibs_dir() == (tmp_path / "pl").resolve()

    def test_backoff(self, tmp_path, monkeypatch):
        monkeypatch.setattr(bootstrap, "pylibs_dir", lambda: tmp_path)
        bootstrap._record_fail(tmp_path, "core", "test")
        assert bootstrap._backoff_active(tmp_path, "core")
        assert not bootstrap._backoff_active(tmp_path, "ocr")
        monkeypatch.setattr(bootstrap, "RETRY_BACKOFF_SECONDS", -1)
        assert not bootstrap._backoff_active(tmp_path, "core")

    def test_core_success_does_not_clear_ocr_backoff(self, tmp_path, monkeypatch):
        bootstrap._record_fail(tmp_path, "ocr", "model unavailable")
        monkeypatch.setattr(bootstrap, "_imports_ok", lambda mods: mods == bootstrap.CORE_IMPORTS)
        assert bootstrap.ensure_core(tmp_path)[0]
        assert bootstrap._backoff_active(tmp_path, "ocr")

    def test_inspect_uses_pylibs_from_data_root(self, tmp_path, monkeypatch, capsys):
        """自检只读，但应和守护预热使用同一 pylibs 导入路径。"""
        target = tmp_path / "pylibs"
        monkeypatch.setattr(bootstrap, "pylibs_dir", lambda: target)
        seen = []
        monkeypatch.setattr(bootstrap, "inject_sys_path", lambda path: seen.append(path))
        monkeypatch.setattr(bootstrap, "_imports_ok", lambda mods: mods == bootstrap.CORE_IMPORTS)
        monkeypatch.setattr(worker, "DATA_DIR", tmp_path)
        monkeypatch.setattr(worker, "STATE_FILE", tmp_path / "state.json")
        assert worker._inspect() == 0
        assert seen == [target]
        assert "core: 就绪" in capsys.readouterr().out

    def test_installer_prefers_local_wheels_without_network(self, tmp_path, monkeypatch):
        wheels = tmp_path / "wheelhouse" / "core"
        wheels.mkdir(parents=True)
        (wheels / "mock-1.0-py3-none-any.whl").write_bytes(b"fake")
        calls = []

        def fake_run(command, **kwargs):
            calls.append(command)
            return subprocess.CompletedProcess(command, 0, "ok", "")

        monkeypatch.setattr(bootstrap.subprocess, "run", fake_run)
        assert bootstrap._install("core", tmp_path / "pylibs") == (True, "ok")
        assert len(calls) == 1
        assert "--no-index" in calls[0] and str(wheels) in calls[0]
        assert "--only-binary=:all:" in calls[0]

    def test_installer_falls_back_to_pypi_after_mirror_failure(self, tmp_path, monkeypatch):
        calls = []

        def fake_run(command, **kwargs):
            calls.append(command)
            return subprocess.CompletedProcess(command, 1 if len(calls) == 1 else 0,
                                               "", "temporary failure")

        monkeypatch.setattr(bootstrap.subprocess, "run", fake_run)
        assert bootstrap._install("ocr", tmp_path / "pylibs") == (True, "ok")
        assert len(calls) == 2
        assert bootstrap.MIRROR_INDEX in calls[0]
        assert bootstrap.PYPI_INDEX in calls[1]

    def test_installer_recovers_pip_error_from_frozen_io(self, tmp_path, monkeypatch):
        io = tmp_path / ".worker-test"
        io.mkdir()
        (io / "pip-stderr").write_text("wheel install error", encoding="utf-8")
        monkeypatch.setenv("CABCLAW_WORKER_IO_DIR", str(io))
        monkeypatch.setattr(bootstrap.sys, "frozen", True, raising=False)
        monkeypatch.setattr(bootstrap.subprocess, "run", lambda command, **kwargs:
                            subprocess.CompletedProcess(command, 1, "", ""))
        ok, error = bootstrap._install("core", tmp_path / "pylibs")
        assert not ok and "wheel install error" in error

    def test_layer_requirement_hashes_do_not_override_each_other(self, tmp_path, monkeypatch):
        installed = []
        monkeypatch.setattr(bootstrap, "_imports_ok", lambda mods: True)
        monkeypatch.setattr(bootstrap, "_install", lambda layer, pylibs: (installed.append(layer) or True, "ok"))
        assert bootstrap.ensure_core(tmp_path)[0]
        assert bootstrap.ensure_ocr(tmp_path)[0]
        assert installed == []
        core_hash = bootstrap._hash_file(tmp_path, "core")
        ocr_hash = bootstrap._hash_file(tmp_path, "ocr")
        assert core_hash.is_file() and ocr_hash.is_file() and core_hash != ocr_hash
        core_hash.write_text("older-core", encoding="utf-8")
        assert bootstrap.ensure_core(tmp_path)[0]
        assert installed == ["core"]
        assert bootstrap.ensure_ocr(tmp_path)[0]
        assert installed == ["core"]

    def test_first_install_injects_new_pylibs_before_import_probe(self, tmp_path, monkeypatch):
        module = "cabclaw_bootstrap_probe"
        pylibs = tmp_path / "fresh-pylibs"
        monkeypatch.setattr(bootstrap, "CORE_IMPORTS", (module,))

        def fake_install(layer, dest):
            dest.mkdir(parents=True)
            (dest / f"{module}.py").write_text("READY = True\n", encoding="utf-8")
            return True, "ok"

        monkeypatch.setattr(bootstrap, "_install", fake_install)
        try:
            assert bootstrap.ensure_core(pylibs) == (True, "ok")
            assert str(pylibs.resolve()) in sys.path
        finally:
            sys.path[:] = [item for item in sys.path if item != str(pylibs.resolve())]
            sys.modules.pop(module, None)

    def test_ocr_engine_passes_string_model_root(self, tmp_path, monkeypatch):
        from types import SimpleNamespace

        seen = []
        fake = SimpleNamespace(__file__=str(tmp_path / "rapidocr" / "__init__.py"),
                               RapidOCR=lambda **kwargs: seen.append(kwargs) or object())
        monkeypatch.setitem(sys.modules, "rapidocr", fake)
        bootstrap.create_ocr_engine()
        assert seen == [{"params": {
            "Global.model_root_dir": str(tmp_path / "rapidocr" / "models")
        }}]


# ---------- parsers ----------

class TestParsers:
    def test_docx(self, samples):
        r = parsers.parse(samples["docx"], worker.DEFAULT_SETTINGS, ocr_ready=False)
        assert r.kind_label == "Word 文档"
        assert "项目计划" in r.titles
        assert "| 阶段 | 时间 |" in r.md
        assert r.total_chars > 0

    def test_xlsx(self, samples):
        r = parsers.parse(samples["xlsx"], worker.DEFAULT_SETTINGS, ocr_ready=False)
        assert "预算" in r.md and "服务器" in r.md

    def test_pptx(self, samples):
        r = parsers.parse(samples["pptx"], worker.DEFAULT_SETTINGS, ocr_ready=False)
        assert "季度汇报" in r.md

    def test_pdf(self, samples):
        r = parsers.parse(samples["pdf"], worker.DEFAULT_SETTINGS, ocr_ready=False)
        assert "Hello Officetools" in r.md
        assert r.kind_label == "PDF 文档"

    def test_xls(self, samples):
        xls = samples_dir() / "sample.xls"
        if not xls.is_file():
            pytest.skip("xls 样例由 dev 脚本生成（xlwt 仅开发环境）")
        r = parsers.parse(xls, worker.DEFAULT_SETTINGS, ocr_ready=False)
        assert "服务器" in r.md

    def test_unsupported(self, samples):
        p = samples["pdf"].with_suffix(".xyz")
        p.write_bytes(b"x")
        with pytest.raises(parsers.ParseError) as ei:
            parsers.parse(p, worker.DEFAULT_SETTINGS, ocr_ready=False)
        assert ei.value.kind == "unsupported"

    def test_md_table_limits(self):
        rows = [["c"] * 50 for _ in range(400)]
        t = parsers._md_table(rows)
        # 表头分隔线 + MAX_CELL_ROWS 行数据（_md_table 保留首行表头 + N 行数据）
        assert t.count("\n") == parsers.MAX_CELL_ROWS + 1
        assert t.splitlines()[0].count("|") == parsers.MAX_CELL_COLS + 1


def samples_dir() -> Path:
    return Path("/tmp/opencode/ot-samples")


# ---------- worker：S1 安全校验 ----------

class TestExtractInboxPath:
    def test_valid(self, env):
        p = env["inbox"] / "a.pdf"
        p.write_bytes(b"x")
        assert worker.extract_inbox_path(f"已存 {p}") == p.resolve()

    def test_traversal_rejected(self, env):
        outside = env["data"] / "evil.pdf"
        outside.parent.mkdir(parents=True, exist_ok=True)
        outside.write_bytes(b"x")
        assert worker.extract_inbox_path(f"已存 {outside}") is None

    def test_relative_rejected(self):
        assert worker.extract_inbox_path("已存 ../../etc/passwd") is None

    def test_resolve_escape_rejected(self, env):
        tricky = f"已存 {env['inbox'] / '..' / 'outside.pdf'}"
        assert worker.extract_inbox_path(tricky) is None

    def test_no_hint(self):
        assert worker.extract_inbox_path("普通消息") is None


# ---------- worker：定位 ----------

class TestLocate:
    def test_point_name(self, env):
        p = env["inbox"] / "合同.pdf"
        p.write_bytes(b"x")
        target, reason = worker.locate_target("解读 合同.pdf", worker.DEFAULT_SETTINGS, False)
        assert target == p and reason == "点名匹配"

    def test_state_fresh_uninterpreted(self, env):
        p = env["inbox"] / "new.pdf"
        p.write_bytes(b"x")
        worker.save_sent_json(worker.STATE_FILE, {
            "latest_file": {"path": str(p), "name": "new.pdf",
                            "received_at": time.time(), "interpreted": False},
        })
        target, _ = worker.locate_target("解读", worker.DEFAULT_SETTINGS, False)
        assert target == p

    def test_state_stale_falls_to_scan(self, env):
        p = env["inbox"] / "old.pdf"
        p.write_bytes(b"x")
        p2 = env["inbox"] / "newer.pdf"
        p2.write_bytes(b"x")
        import os
        os.utime(p, (time.time() - 100, time.time() - 100))
        worker.save_sent_json(worker.STATE_FILE, {
            "latest_file": {"path": str(p), "name": "old.pdf",
                            "received_at": time.time() - 25 * 3600, "interpreted": False},
        })
        target, _ = worker.locate_target("解读", worker.DEFAULT_SETTINGS, False)
        assert target == p2

    def test_none(self, env):
        target, _ = worker.locate_target("解读", worker.DEFAULT_SETTINGS, False)
        assert target is None


# ---------- worker：入站流程 ----------

class TestInboundFlow:
    def test_file_msg_default_rc3(self, env):
        p = env["inbox"] / "f.pdf"
        p.write_bytes(b"x")
        rc, _ = worker.handle_inbound(f"[收到file: f.pdf，已存 {p}]", "c")
        assert rc == 3
        state = json.loads(worker.STATE_FILE.read_text())
        assert state["latest_file"]["interpreted"] is False

    def test_file_msg_dependency_unready_and_auto_off_returns_rc3(self, env, monkeypatch):
        """依赖未就绪时，自动解析关闭仍必须记录并交还 Agent。"""
        monkeypatch.setattr(worker.bootstrap, "ensure_core", lambda: (False, "未安装"))
        p = env["inbox"] / "f.pdf"
        p.write_bytes(b"x")
        settings = dict(worker.DEFAULT_SETTINGS, file_auto_on=False)
        monkeypatch.setattr(worker, "_settings", lambda: settings)
        rc, _ = worker.handle_inbound(f"[收到file: f.pdf，已存 {p}]", "c")
        assert rc == 3
        assert json.loads(worker.STATE_FILE.read_text())["latest_file"]["interpreted"] is False

    def test_image_system_marker_is_recorded_and_returned(self, env, monkeypatch):
        monkeypatch.setattr(worker.bootstrap, "ensure_core", lambda: (False, "未安装"))
        p = env["inbox"] / "photo.png"
        p.write_bytes(b"\x89PNG\r\n\x1a\n")
        settings = dict(worker.DEFAULT_SETTINGS, file_auto_on=False)
        monkeypatch.setattr(worker, "_settings", lambda: settings)
        rc, _ = worker.handle_inbound(f"[收到image: photo.png，已存 {p}]", "c")
        assert rc == 3
        assert json.loads(worker.STATE_FILE.read_text())["latest_file"]["name"] == "photo.png"

    def test_file_msg_auto_on(self, env, samples, monkeypatch):
        monkeypatch.setattr(worker, "_ocr_ready", lambda s: False)
        settings = dict(worker.DEFAULT_SETTINGS, file_auto_on=True)
        monkeypatch.setattr(worker, "_settings", lambda: settings)
        p = env["inbox"] / "sample.pdf"
        p.write_bytes(samples["pdf"].read_bytes())
        rc, reply = worker.handle_inbound(f"[收到file: sample.pdf，已存 {p}]", "c")
        assert rc == 0 and "文档解读" in reply
        state = json.loads(worker.STATE_FILE.read_text())
        assert state["latest_file"]["interpreted"] is True
        assert (Path(state["latest_output"])).is_file()

    def test_interpret_flow(self, env, samples):
        p = env["inbox"] / "sample.pdf"
        p.write_bytes(samples["pdf"].read_bytes())
        rc, reply = worker.handle_inbound("解读", "c")
        assert rc == 0 and "PDF 文档" in reply
        # 二次解读：已解读 → 无新文件
        rc2, reply2 = worker.handle_inbound("解读", "c")
        assert rc2 == 0 and "没有找到" in reply2

    def test_dry_run_no_side_effect(self, env, samples):
        p = env["inbox"] / "sample.pdf"
        p.write_bytes(samples["pdf"].read_bytes())
        worker.handle_inbound(f"[收到file: sample.pdf，已存 {p}]", "c")
        rc, reply = worker.handle_inbound("解读", "c", dry=True)
        assert rc == 0 and reply.startswith("[dry]")
        assert not json.loads(worker.STATE_FILE.read_text())["latest_file"]["interpreted"]

    def test_range_ext_rc0(self, env):
        p = env["inbox"] / "v.mp4"
        p.write_bytes(b"x")
        rc, reply = worker.handle_inbound(f"[收到file: v.mp4，已存 {p}]", "c", dry=False)
        assert rc == 0 and "暂不支持" in reply

    def test_point_missing(self, env):
        rc, reply = worker.handle_inbound("解读 不存在.pdf", "c")
        assert rc == 0 and "没有名为" in reply

    def test_parse_lock_is_nonblocking_and_released(self, env):
        lock_path = worker.LOCK_FILE
        lock_path.parent.mkdir(parents=True, exist_ok=True)
        with lock_path.open("a+b") as first, lock_path.open("a+b") as second:
            assert worker._acquire_parse_lock(first)
            assert not worker._acquire_parse_lock(second)
            worker._release_parse_lock(first)
            assert worker._acquire_parse_lock(second)
            worker._release_parse_lock(second)

    def test_windows_parse_lock_uses_msvcrt(self, env, monkeypatch):
        """离线仿真 Windows 文件锁分支，真实 Windows 结果留待远端测试。"""
        import errno
        from types import SimpleNamespace

        held = set()

        def locking(fd, mode, length):
            if mode == 1:  # LK_UNLCK
                held.remove(0)
            elif held:
                raise OSError(errno.EACCES, "locked")
            else:
                held.add(0)

        monkeypatch.setattr(worker, "os", SimpleNamespace(name="nt", SEEK_END=os.SEEK_END))
        monkeypatch.setattr(worker, "msvcrt", SimpleNamespace(LK_NBLCK=2, LK_UNLCK=1,
                                                               locking=locking), raising=False)
        path = env["data"] / ".parse.lock"
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("a+b") as first, path.open("a+b") as second:
            assert worker._acquire_parse_lock(first)
            assert not worker._acquire_parse_lock(second)
            worker._release_parse_lock(first)
            assert worker._acquire_parse_lock(second)
            worker._release_parse_lock(second)

    def test_cleanup_retention(self, env, monkeypatch):
        outputs = env["data"] / "outputs"
        outputs.mkdir(parents=True)
        old = outputs / "old-20260101-000000.md"
        old.write_text("x")
        import os
        os.utime(old, (time.time() - 8 * 86400, time.time() - 8 * 86400))
        fresh = outputs / "fresh-20260904-000000.md"
        fresh.write_text("x")
        worker._cleanup_outputs(7)
        assert not old.exists() and fresh.exists()

    def test_overview_truncate(self, env, samples, monkeypatch):
        monkeypatch.setattr(worker, "_ocr_ready", lambda s: False)
        settings = dict(worker.DEFAULT_SETTINGS, overview_max_chars=100)
        rc, reply = worker.parse_and_deliver(
            samples["pdf"], settings, dry=False
        ) if False else (0, "")
        # 直接测 _overview_text 截断
        r = parsers.parse(samples["pdf"], settings, ocr_ready=False)
        text = worker._overview_text(r, "/tmp/x.md", settings)
        assert len(text) <= 100 + len("…（已截断）")


# ---------- 部署形态：裸 spawn（无 PYTHONPATH 注入，对齐 emotion 先例） ----------

def test_bare_spawn_no_pythonpath(tmp_path):
    """模拟部署裸 spawn：bridge inbound（main.py:368）不注入 PYTHONPATH，
    worker 必须自持 sys.path（项目根 bridge + modules/ common）才能 import 成功。

    搭临时部署树（symlink 宿主 bridge/common + 拷贝模块），env 剥掉 PYTHONPATH、
    cwd 挪走——修复前此用例必红（ImportError: No module named 'bridge'）。
    --inspect 零副作用（不装依赖不写文件）。
    """
    import shutil

    host = Path(os.environ["CABCLAW_HOST"]).resolve()
    deploy = tmp_path / "deploy"
    (deploy / "modules").mkdir(parents=True)
    (deploy / "bridge").symlink_to(host / "bridge", target_is_directory=True)
    (deploy / "modules" / "common").symlink_to(host / "modules" / "common", target_is_directory=True)
    shutil.copytree(MODULE_DIR, deploy / "modules" / "Officetools",
                    ignore=shutil.ignore_patterns("__pycache__"))
    env = {k: v for k, v in os.environ.items() if k != "PYTHONPATH"}
    proc = subprocess.run(
        [sys.executable, str(deploy / "modules" / "Officetools" / "officetools_worker.py"), "--inspect"],
        capture_output=True, text=True, env=env,
        cwd=str(tmp_path),  # cwd 也挪走：证明不依赖 cwd
        timeout=60,
    )
    assert proc.returncode == 0, f"裸 spawn 失败:\n{proc.stderr[-800:]}"
    assert "state:" in proc.stdout
    assert "core:" in proc.stdout
